import os
import traceback
from flask import Flask, request, send_file
from flask_cors import CORS
from pdf2docx import Converter
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

app = Flask(__name__)
CORS(app)

# 100% محفوظ پرانا روٹ
@app.route('/convert', methods=['POST'])
def convert_pdf():
    if 'file' not in request.files:
        return 'No file uploaded', 400
    
    file = request.files['file']
    pdf_path = 'temp.pdf'
    docx_path = 'temp.docx'
    
    file.save(pdf_path)
    
    cv = Converter(pdf_path)
    cv.convert(docx_path, start=0, end=None)
    cv.close()
    
    return send_file(docx_path, as_attachment=True)

# ورڈ سے پی ڈی ایف (فورسڈ ٹیبل پارسنگ)
@app.route('/convert-word', methods=['POST'])
def convert_word_to_pdf():
    try:
        if 'file' not in request.files:
            return 'No file uploaded', 400
        
        file = request.files['file']
        filename_base = os.path.splitext(file.filename)[0]
        docx_path = f"temp_{filename_base}.docx"
        pdf_path = f"temp_{filename_base}.pdf"
        
        file.save(docx_path)
        
        doc = Document(docx_path)
        pdf_doc = SimpleDocTemplate(
            pdf_path, 
            pagesize=letter,
            rightMargin=20, leftMargin=20,
            topMargin=20, bottomMargin=20
        )
        
        styles = getSampleStyleSheet()
        normal_style = styles['Normal']
        normal_style.fontSize = 8
        normal_style.leading = 10

        cell_style = ParagraphStyle(
            'CellStyle',
            parent=normal_style,
            fontSize=7.5,
            leading=9
        )
        
        story = []
        
        # ورڈ کے تمام ٹیبلز اور پیراگراف کو نکالنے کا حتمی طریقہ
        for table in doc.tables:
            table_data = []
            for row in table.rows:
                row_data = []
                for cell in row.cells:
                    cell_text = "\n".join([p.text.strip() for p in cell.paragraphs if p.text.strip()])
                    row_data.append(Paragraph(cell_text, cell_style))
                table_data.append(row_data)
            
            if table_data:
                t = Table(table_data)
                t.setStyle(TableStyle([
                    ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#2D3748')),
                    ('INNERGRID', (0,0), (-1,-1), 0.25, colors.HexColor('#CBD5E0')),
                    ('VALIGN', (0,0), (-1,-1), 'TOP'),
                    ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#EDF2F7')),
                    ('BOTTOMPADDING', (0,0), (-1,-1), 3),
                    ('TOPPADDING', (0,0), (-1,-1), 3),
                    ('LEFTPADDING', (0,0), (-1,-1), 3),
                    ('RIGHTPADDING', (0,0), (-1,-1), 3),
                ]))
                story.append(t)
                story.append(Spacer(1, 8))

        # اگر فائل میں الگ سے بھی پیراگراف ہوں
        for p in doc.paragraphs:
            if p.text.strip():
                story.append(Paragraph(p.text, normal_style))
                story.append(Spacer(1, 4))
            
        pdf_doc.build(story)
        
        if os.path.exists(pdf_path):
            response = send_file(pdf_path, as_attachment=True, download_name=f"{filename_base}.pdf")
            return response
        else:
            return "PDF generation failed", 500
            
    except Exception as e:
        error_details = traceback.format_exc()
        print(error_details)
        return f"Conversion Error: {str(e)}\n{error_details}", 500

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 5000)))
